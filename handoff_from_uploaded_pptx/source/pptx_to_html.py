import sys, html, re
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE
from lxml import etree
NS={'a':'http://schemas.openxmlformats.org/drawingml/2006/main','p':'http://schemas.openxmlformats.org/presentationml/2006/main'}
PX=96/914400
def px(v): return round((v or 0)*PX,2)
def color_of(el):
    if el is None: return None,1
    c=el.find('a:srgbClr',NS)
    if c is None: return None,1
    a=c.find('a:alpha',NS)
    return '#'+c.get('val'), (int(a.get('val'))/100000 if a is not None else 1)
def fill_css(sp):
    spPr=sp._element.find('.//p:spPr',NS)
    if spPr is None: return ''
    if spPr.find('a:noFill',NS) is not None: return 'background:transparent;'
    sf=spPr.find('a:solidFill',NS)
    col,al=color_of(sf)
    if col:
        r,g,b=int(col[1:3],16),int(col[3:5],16),int(col[5:7],16)
        return f'background:rgba({r},{g},{b},{al});'
    return ''
def line_css(sp):
    spPr=sp._element.find('.//p:spPr',NS)
    ln=spPr.find('a:ln',NS) if spPr is not None else None
    if ln is None: return ''
    if ln.find('a:noFill',NS) is not None: return 'border:none;'
    col,al=color_of(ln.find('a:solidFill',NS))
    if not col: return ''
    w=int(ln.get('w','12700'))/12700*96/72
    r,g,b=int(col[1:3],16),int(col[3:5],16),int(col[5:7],16)
    d=ln.find('a:prstDash',NS); st='dashed' if d is not None and d.get('val') not in('solid',) else 'solid'
    return f'border:{w:.2f}px {st} rgba({r},{g},{b},{al});'
def text_html(tf_el):
    out=[]
    for p in tf_el.findall('a:p',NS):
        pPr=p.find('a:pPr',NS); al={'ctr':'center','r':'right','l':'left'}.get(pPr.get('algn') if pPr is not None else None,'left')
        runs=[]
        for r in p.findall('a:r',NS):
            rPr=r.find('a:rPr',NS); st=[]
            if rPr is not None:
                if rPr.get('sz'): st.append(f"font-size:{int(rPr.get('sz'))/100*96/72:.2f}px")
                if rPr.get('b')=='1': st.append('font-weight:700')
                if rPr.get('i')=='1': st.append('font-style:italic')
                c,_=color_of(rPr.find('a:solidFill',NS))
                if c: st.append(f'color:{c}')
                lat=rPr.find('a:latin',NS)
                if lat is not None: st.append(f"font-family:'{lat.get('typeface')}',Calibri,Arial,sans-serif")
                if rPr.get('spc'): st.append(f"letter-spacing:{int(rPr.get('spc'))/100*96/72:.2f}px")
            runs.append(f'<span style="{";".join(st)}">{html.escape(r.findtext("a:t",namespaces=NS) or "")}</span>')
        for br in p.findall('a:br',NS): runs.append('<br>')
        out.append(f'<div style="text-align:{al}">{"".join(runs) or "&nbsp;"}</div>')
    return ''.join(out)
def shape_html(sh):
    x,y,w,h=px(sh.left),px(sh.top),px(sh.width),px(sh.height)
    rot=getattr(sh,'rotation',0) or 0
    base=f'position:absolute;left:{x}px;top:{y}px;width:{w}px;height:{h}px;'
    if rot: base+=f'transform:rotate({rot}deg);'
    el=sh._element
    if el.tag.endswith('}cxnSp') or (el.tag.endswith('}sp') and el.find('.//a:prstGeom',NS) is not None and el.find('.//a:prstGeom',NS).get('prst')=='line'):
        ln=el.find('.//a:ln',NS); col,al=color_of(ln.find('a:solidFill',NS)) if ln is not None else (None,1)
        col=col or '#888'; lw=int(ln.get('w','12700'))/12700*96/72 if ln is not None else 1
        d=ln.find('a:prstDash',NS) if ln is not None else None
        dash={'dash':'6 4','sysDash':'3 2','sysDot':'1 2','dot':'1 3'}.get(d.get('val') if d is not None else None,'')
        xf=el.find('.//a:xfrm',NS); fh=xf.get('flipH')=='1'; fv=xf.get('flipV')=='1'
        x1,x2=(w,0) if fh else (0,w); y1,y2=(h,0) if fv else (0,h)
        return f'<svg style="position:absolute;left:{x}px;top:{y}px;overflow:visible" width="{max(w,1)}" height="{max(h,1)}"><line x1="{x1}" y1="{y1}" x2="{x2}" y2="{y2}" stroke="{col}" stroke-opacity="{al}" stroke-width="{lw:.2f}" stroke-dasharray="{dash}"/></svg>'
    if sh.has_text_frame is False and sh.shape_type!=MSO_SHAPE_TYPE.PICTURE and not getattr(sh,'has_table',False):
        return ''
    if getattr(sh,'has_table',False) and sh.has_table:
        rows=[]
        for r in sh.table.rows:
            tds=[]
            for c in r.cells:
                tcPr=c._tc.find('a:tcPr',NS); f=color_of(tcPr.find('a:solidFill',NS)) if tcPr is not None else (None,1)
                bg=f'background:{f[0]};' if f[0] else ''
                tds.append(f'<td style="{bg}padding:4px 8px;vertical-align:middle">{text_html(c._tc.find("a:txBody",NS))}</td>')
            rows.append(f'<tr style="height:{px(r.height)}px">{"".join(tds)}</tr>')
        cols=''.join(f'<col style="width:{px(c.width)}px">' for c in sh.table.columns)
        return f'<table style="{base}border-collapse:collapse;table-layout:fixed">{cols}{"".join(rows)}</table>'
    if sh.shape_type==MSO_SHAPE_TYPE.PICTURE:
        import base64
        b=base64.b64encode(sh.image.blob).decode()
        return f'<img style="{base}" src="data:{sh.image.content_type};base64,{b}">'
    prst=el.find('.//a:prstGeom',NS); pr=prst.get('prst') if prst is not None else 'rect'
    css=base+fill_css(sh)+line_css(sh)+'box-sizing:border-box;'
    if pr=='ellipse': css+='border-radius:50%;'
    if pr=='roundRect':
        adj=el.find('.//a:gd',NS); v=int(adj.get('fmla').split()[1])/100000 if adj is not None else .16667
        css+=f'border-radius:{min(w,h)*v:.1f}px;'
    tx=el.find('.//p:txBody',NS)
    inner=''
    if tx is not None and ''.join(tx.itertext()).strip():
        bp=tx.find('a:bodyPr',NS)
        ins={k:int(bp.get(k,d))*PX for k,d in (('lIns',91440),('tIns',45720),('rIns',91440),('bIns',45720))} if bp is not None else {'lIns':9,'tIns':4,'rIns':9,'bIns':4}
        anchor={'ctr':'center','b':'flex-end'}.get(bp.get('anchor') if bp is not None else None,'flex-start')
        css+=f'display:flex;flex-direction:column;justify-content:{anchor};padding:{ins["tIns"]:.1f}px {ins["rIns"]:.1f}px {ins["bIns"]:.1f}px {ins["lIns"]:.1f}px;line-height:1.15;'
        inner=text_html(tx)
    return f'<div style="{css}">{inner}</div>'
prs=Presentation(sys.argv[1]); out=sys.argv[2]
W,H=px(prs.slide_width),px(prs.slide_height)
titles=[];slides=[]
for i,s in enumerate(prs.slides,1):
    bg='#0D1B2A'
    bgel=s._element.find('.//p:bg//a:srgbClr',NS)
    if bgel is not None: bg='#'+bgel.get('val')
    body=''.join(shape_html(sh) for sh in s.shapes)
    t=next((sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()),f'Slide {i}')
    titles.append(t.split('\n')[0]); slides.append((bg,body))
css=f"*{{box-sizing:border-box}}body{{margin:0;background:#050b12;font-family:Calibri,Arial,sans-serif;color:#E8F4FF}}.slide{{position:relative;width:{W}px;height:{H}px;margin:0 auto;overflow:hidden}}.slide div{{white-space:normal}}.tabs{{text-align:center;padding:12px;position:sticky;top:0;background:#050b12;z-index:9}}.tabs button{{background:#0F1F30;color:#7AAAC5;border:1px solid #1E3A52;padding:6px 16px;margin:0 6px;border-radius:6px;cursor:pointer;font-size:13px}}.tabs button.on{{background:#7C3AED;color:#fff}}.pg{{display:none}}.pg.on{{display:block}}"
btn=''.join(f'<button data-i="{i}" class="{"on" if i==0 else ""}">Slide {i+1}</button>' for i in range(len(slides)))
pgs=''.join(f'<div class="pg{" on" if i==0 else ""}"><div class="slide" style="background:{bg}">{b}</div></div>' for i,(bg,b) in enumerate(slides))
open(out,'w',encoding='utf-8').write(f'<!doctype html><html><head><meta charset="utf-8"><title>FoxBrain Benchmark Priority Radar</title><style>{css}</style></head><body><div class="tabs">{btn}</div>{pgs}<script>document.querySelectorAll(".tabs button").forEach(b=>b.onclick=()=>{{document.querySelectorAll(".tabs button,.pg").forEach(e=>e.classList.remove("on"));b.classList.add("on");document.querySelectorAll(".pg")[b.dataset.i].classList.add("on")}})</script></body></html>')
# per-slide files too
for i,(bg,b) in enumerate(slides,1):
    open(out.replace('all_slides',f'slide{i}'),'w',encoding='utf-8').write(f'<!doctype html><html><head><meta charset="utf-8"><title>{html.escape(titles[i-1])}</title><style>{css}.pg{{display:block}}</style></head><body><div class="pg"><div class="slide" style="background:{bg}">{b}</div></div></body></html>')
print('ok',len(slides))
